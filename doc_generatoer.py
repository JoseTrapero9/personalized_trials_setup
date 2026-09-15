import csv
import glob
import os
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


def set_cell_background(cell, hex_color):
    """Applies background shading to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def natural_sort_key(filename):
    """Sorts filenames naturally (e.g., s01 before s02)."""
    parts = re.split(r"(\d+)", os.path.basename(filename))
    return [int(p) if p.isdigit() else p.lower() for p in parts]


def convert_logs_to_word(
    input_dir="log_files", output_file="Experiment_Logs.docx"
):
    doc = Document()

    # Document Header
    title = doc.add_heading("Experiment Log Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Compiled trial data across subjects and conditions.")

    files = sorted(
        glob.glob(os.path.join(input_dir, "*.txt")), key=natural_sort_key
    )
    if not files:
        print(f"No log files found in '{input_dir}'.")
        return

    current_subject = None

    for filepath in files:
        fname = os.path.basename(filepath)
        # Parse subject ID and condition name from filename (e.g., s01_baseline.txt)
        match = re.match(r"(s\d+)_(.+)\.txt", fname)
        subj_tag, cond_name = (
            match.groups() if match else ("General", fname.replace(".txt", ""))
        )

        # Start a new page for every new subject
        if subj_tag != current_subject:
            if current_subject is not None:
                doc.add_page_break()
            current_subject = subj_tag
            doc.add_heading(f"Subject: {subj_tag.upper()}", level=1)

        # Condition subheading
        doc.add_heading(
            f"Condition: {cond_name.replace('_', ' ').title()}", level=2
        )

        # Read TSV data
        with open(filepath, "r", encoding="utf-8") as f:
            reader = list(csv.reader(f, delimiter="\t"))
            if not reader:
                continue
            headers = [h.title() for h in reader[0]]
            rows = reader[1:]

        # Build table
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Format header row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            set_cell_background(hdr_cells[i], "2F4F4F")  # Dark slate gray
            run = hdr_cells[i].paragraphs[0].runs[0]
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

        # Populate rows
        for row_idx, row_data in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            is_even = row_idx % 2 == 0
            for col_idx, cell_value in enumerate(row_data):
                row_cells[col_idx].text = cell_value
                run = row_cells[col_idx].paragraphs[0].runs[0]
                run.font.size = Pt(8.5)
                # Alternating row tint
                if is_even:
                    set_cell_background(row_cells[col_idx], "F7F9FA")

        # Set column widths
        col_widths = [
            Inches(0.8),
            Inches(0.8),
            Inches(2.5),
            Inches(2.4),
        ]  # Piece, Sum, Instr, Answer
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width

        doc.add_paragraph()  # Spacing

    doc.save(output_file)
    print(f"Successfully generated: {output_file}")


if __name__ == "__main__":
    convert_logs_to_word()